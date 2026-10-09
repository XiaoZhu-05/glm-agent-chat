# -*- coding: utf-8 -*-
"""Excel 提取（服务端调用）：python extract_xlsx.py <file> [max_rows]
输出单行 JSON：{ok, sheets:[{name, rows, cols, truncated, table(markdown)}], error?}
说明：data_only=True → 公式单元格返回缓存的计算值；合并单元格仅左上角单元格有值，其余为空。
"""
import sys
import json

try:
    import openpyxl
except ImportError:
    print(json.dumps({"ok": False, "error": "服务器缺少 openpyxl 库（pip install openpyxl）"}, ensure_ascii=False))
    sys.exit(0)


def cell(v):
    if v is None:
        return ""
    return str(v)


def main():
    path = sys.argv[1]
    max_rows = int(sys.argv[2]) if len(sys.argv) > 2 else 200
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows_data = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= max_rows:
                    break
                rows_data.append([cell(c) for c in row])
            # markdown 表格（首行作为表头）
            lines = []
            if rows_data:
                width = max(len(r) for r in rows_data)
                for r in rows_data:
                    r = r + [""] * (width - len(r))
                    lines.append("| " + " | ".join(x.replace("|", "\\|").replace("\n", " ") for x in r) + " |")
                sep = "| " + " | ".join(["---"] * width) + " |"
                lines.insert(1, sep)
            sheets.append({
                "name": ws.title,
                "rows": ws.max_row or len(rows_data),
                "cols": ws.max_column or 0,
                "truncated": (ws.max_row or 0) > max_rows,
                "table": "\n".join(lines) if lines else "（空工作表）",
            })
        print(json.dumps({"ok": True, "sheets": sheets}, ensure_ascii=False))
    except Exception as e:
        msg = str(e)
        if "Excel 2007" in msg or ".xls" in path.lower().endswith(".xls"):
            msg = "暂不支持旧版 .xls 二进制格式，请在 Excel 中另存为 .xlsx 后重新上传"
        print(json.dumps({"ok": False, "error": f"Excel 解析失败：{msg}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
